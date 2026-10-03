from __future__ import annotations

import sys
from pathlib import Path

from ytcore.content.dokum import DokumBolum
from ytcore.models import IndexKaydi
from ytcore.output.klasor import IMZA
from ytcore.output.typst_runtime import compile_pdf


def _zaman(sn: int | None) -> str:
    if sn is None:
        return ""
    return f"[{sn // 60:02d}:{sn % 60:02d}] "


def _md(kayit: IndexKaydi, bolumler: list[DokumBolum]) -> str:
    sure = f"{kayit.sure_sn} sn" if kayit.sure_sn else "?"
    url = kayit.kaynak_url or kayit.video_url
    sahip_etiketi = "Kanal" if kayit.kaynak_turu == "youtube" else "Sahip/Yazar"
    sahip = kayit.kaynak_sahibi or kayit.kanal or "?"
    dokum_adi = "Kronolojik Döküm" if kayit.kaynak_turu == "youtube" else "Yapısal Döküm"
    sure_bilgisi = f" · Süre: {sure}" if kayit.kaynak_turu == "youtube" else ""
    bas = (
        f"# Rasathane · {kayit.baslik} — {dokum_adi}\n\n"
        f"> Kaynak: {url}\n"
        f"> {sahip_etiketi}: {sahip}{sure_bilgisi} · Dil: {kayit.anadil or '?'} · "
        f"Analiz: {kayit.analiz_tarihi}\n\n"
    )
    govde = []
    for b in bolumler:
        kw = f"\n\n**Anahtar:** {', '.join(b.keywords)}" if b.keywords else ""
        govde.append(f"## {_zaman(b.baslangic_sn)}{b.baslik}\n\n{b.metin}{kw}")
    return bas + "\n\n".join(govde) + f"\n\n---\n_{IMZA}_\n"


def dokum_yaz(
    klasor: Path,
    kayit: IndexKaydi,
    bolumler: list[DokumBolum],
    *,
    pdf: bool = True,
    docx: bool = True,
) -> list[Path]:
    """03_dokum.md (her zaman) + opsiyonel .docx (python-docx) + .pdf (typst). md birincil;
    pdf/docx best-effort — patlasa md kaybolmaz (uyarı stderr; Faz 1 'md kaybetme' dersi)."""
    paths: list[Path] = []
    md_metin = _md(kayit, bolumler)
    md_p = klasor / "03_dokum.md"
    md_p.write_text(md_metin, encoding="utf-8")
    paths.append(md_p)
    if pdf:
        try:
            paths.append(_pdf_yaz(klasor, md_metin))
        except Exception as e:  # noqa: BLE001 — pdf best-effort, md birincil
            print(f"[uyari] PDF uretilemedi (md birincil): {e}", file=sys.stderr)
    if docx:
        try:
            paths.append(_docx_yaz(klasor, kayit, bolumler))
        except Exception as e:  # noqa: BLE001 — docx best-effort
            print(f"[uyari] DOCX uretilemedi (md birincil): {e}", file=sys.stderr)
    return paths


def _pdf_yaz(klasor: Path, md_metin: str) -> Path:
    """Typst ile PDF (md'yi raw blok olarak göm). Typst yoksa FileNotFoundError → best-effort.
    Profesyonel şablon Faz 5 (1350+ Typst Universe şablonu)."""
    typ = klasor / "_dokum.typ"
    kacis = md_metin.replace("\\", "\\\\").replace('"', '\\"')
    typ.write_text(
        '#set text(font: ("Arial", "DejaVu Sans"), size: 11pt)\n#set page(margin: 2cm)\n'
        f'#raw("{kacis}", block: true, lang: "markdown")\n',
        encoding="utf-8",
    )
    pdf_p = klasor / "03_dokum.pdf"
    try:
        compile_pdf(typ, pdf_p)
    finally:
        # Typst yoksa/derleme patlasa da geçici .typ kullanıcının teslim klasöründe KALMASIN
        # (review MED #14: typst'siz makinede her döküm yazımında _dokum.typ çöpü sızıyordu).
        typ.unlink(missing_ok=True)
    return pdf_p


def _docx_yaz(klasor: Path, kayit: IndexKaydi, bolumler: list[DokumBolum]) -> Path:
    """python-docx ile programatik DOCX (docxtpl şablon-tabanlı yol Faz 5; programatik
    çok-satır XML Jinja tuzağından kaçınır — A03 RichText uyarısı)."""
    from docx import Document

    doc = Document()
    dokum_adi = "Kronolojik Döküm" if kayit.kaynak_turu == "youtube" else "Yapısal Döküm"
    sahip_etiketi = "Kanal" if kayit.kaynak_turu == "youtube" else "Sahip/Yazar"
    sahip = kayit.kaynak_sahibi or kayit.kanal or "?"
    doc.add_heading(f"Rasathane · {kayit.baslik} — {dokum_adi}", level=0)
    doc.add_paragraph(f"Kaynak: {kayit.kaynak_url or kayit.video_url} · {sahip_etiketi}: {sahip}")
    for b in bolumler:
        doc.add_heading(f"{_zaman(b.baslangic_sn)}{b.baslik}".strip(), level=1)
        doc.add_paragraph(b.metin)
        if b.keywords:
            doc.add_paragraph("Anahtar: " + ", ".join(b.keywords))
    doc.add_paragraph(f"\n{IMZA}")
    p = klasor / "03_dokum.docx"
    doc.save(str(p))
    return p
